"""
DOCX extraction service.

Uses python-docx to extract paragraphs, native Word tables, and
embedded images in their actual document order.

Design decisions:
- Iterates document body XML children to preserve paragraph/table order.
- Extracts embedded image blobs from the docx package and saves them as assets.
- Floating shapes and text boxes are reported as limitations.
- Table rows are read cell-by-cell; merged cells carry the spanned text.
- Images embedded in paragraph runs are extracted and stored.
"""

from __future__ import annotations

import io
import logging
import os
import re
import uuid
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from backend.schema import (
    BlockType, ContentBlock, DocumentResult, ExtractionMethod,
    Geometry, ProcessingSummary, SourceFormat, TableCell,
    TableData, TableRow, ValidationStatus,
)
from backend.services.table_reconstructor import TableReconstructor
from backend.services.image_table_recognizer import ImageTableRecognizer

logger = logging.getLogger(__name__)

reconstructor = TableReconstructor()


class DocxExtractor:
    """Extract a DOCX file into a DocumentResult."""

    def __init__(self, assets_folder: str):
        self.assets_folder = assets_folder
        os.makedirs(assets_folder, exist_ok=True)
        self.table_recognizer = ImageTableRecognizer()

    # ──────────────────────────────────────────────────────────────
    # Public API
    # ──────────────────────────────────────────────────────────────

    def extract(self, docx_path: str, document_id: str, filename: str) -> DocumentResult:
        result = DocumentResult(
            document_id=document_id,
            filename=filename,
            source_format=SourceFormat.DOCX,
        )
        result.summary.total_pages = 1  # DOCX has no strict page concept at library level

        try:
            self._extract_all(docx_path, document_id, result)
        except Exception as exc:
            logger.exception("DOCX extraction failed for %s", docx_path)
            result.summary.warnings.append(f"Extraction error: {exc}")

        result.compute_summary()
        self._check_unsupported_elements(docx_path, result)
        return result

    # ──────────────────────────────────────────────────────────────
    # Internal helpers
    # ──────────────────────────────────────────────────────────────

    def _extract_all(self, docx_path: str, document_id: str, result: DocumentResult) -> None:
        doc = Document(docx_path)
        reading_order = 0
        section_index = 1   # logical "page" approximation for DOCX

        # Build an image relationship map: rId -> image bytes + extension
        image_map = self._build_image_map(doc)

        # Iterate body children in XML order to preserve paragraph/table interleaving
        body = doc.element.body
        for child in body:
            tag = child.tag

            if tag == qn("w:p"):
                # Paragraph
                para = Paragraph(child, doc)
                text = para.text.strip()

                # Check for embedded images in runs
                image_blocks = self._extract_images_from_paragraph(
                    child, image_map, document_id, section_index, reading_order
                )
                for ib in image_blocks:
                    if ib.asset_id:
                        asset_path = os.path.join(self.assets_folder, ib.asset_id)
                        if os.path.isfile(asset_path):
                            tbl_res = self.table_recognizer.recognize_table(asset_path, ib.asset_id)
                            if tbl_res.success and tbl_res.table_data:
                                ib.block_type = BlockType.TABLE
                                ib.extraction_method = ExtractionMethod.OCR_TABLE
                                ib.table_data = tbl_res.table_data
                                ib.validation_status = tbl_res.validation_status
                                ib.metadata["is_reconstructed_table"] = True
                                ib.metadata["original_asset_id"] = ib.asset_id
                                ib.notes.append("Reconstructed from embedded image table via local OCR.")
                    result.blocks.append(ib)
                    reading_order += 1

                if text:
                    block = ContentBlock(
                        block_type=BlockType.TEXT,
                        source_format=SourceFormat.DOCX,
                        page_number=section_index,
                        reading_order=reading_order,
                        extraction_method=ExtractionMethod.NATIVE_TEXT,
                        validation_status=ValidationStatus.VALID,
                        text=text,
                        metadata={
                            "style": para.style.name if para.style else "",
                        },
                    )
                    result.blocks.append(block)
                    reading_order += 1

                # Detect page breaks
                if self._has_page_break(child):
                    section_index += 1

            elif tag == qn("w:tbl"):
                # Table
                table = Table(child, doc)
                table_data = self._extract_table(table)
                if table_data:
                    table_data.html = reconstructor.to_html(table_data)
                    validation = self._validate_table(table_data)
                    block = ContentBlock(
                        block_type=BlockType.TABLE,
                        source_format=SourceFormat.DOCX,
                        page_number=section_index,
                        reading_order=reading_order,
                        extraction_method=ExtractionMethod.NATIVE_TABLE,
                        validation_status=validation,
                        table_data=table_data,
                    )
                    result.blocks.append(block)
                    reading_order += 1

            elif tag == qn("w:sectPr"):
                section_index += 1

        result.summary.total_pages = section_index

    # ── Image handling ────────────────────────────────────────────

    def _build_image_map(self, doc: Document) -> dict[str, tuple[bytes, str]]:
        """
        Returns {rId: (image_bytes, extension)} for all images in the docx package.
        """
        image_map: dict[str, tuple[bytes, str]] = {}
        try:
            for rel_id, rel in doc.part.rels.items():
                if "image" in rel.reltype:
                    try:
                        img_part   = rel.target_part
                        img_bytes  = img_part.blob
                        content_type = img_part.content_type   # e.g. image/png
                        ext = content_type.split("/")[-1].lower()
                        if ext in ("jpeg", "jpg"):
                            ext = "jpg"
                        image_map[rel_id] = (img_bytes, ext)
                    except Exception as exc:
                        logger.debug("Could not load image rId=%s: %s", rel_id, exc)
        except Exception as exc:
            logger.warning("Could not build image map: %s", exc)
        return image_map

    def _extract_images_from_paragraph(
        self,
        para_elem: Any,
        image_map: dict[str, tuple[bytes, str]],
        document_id: str,
        section_index: int,
        reading_order: int,
    ) -> list[ContentBlock]:
        """Find <a:blip r:embed="rIdX"/> references in a paragraph."""
        blocks: list[ContentBlock] = []
        ns = {
            "a":   "http://schemas.openxmlformats.org/drawingml/2006/main",
            "r":   "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
        }
        for blip in para_elem.iter(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
        ):
            r_embed = blip.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
            )
            if r_embed and r_embed in image_map:
                img_bytes, ext = image_map[r_embed]
                asset_filename = f"{document_id}_s{section_index}_{r_embed}.{ext}"
                asset_path     = os.path.join(self.assets_folder, asset_filename)
                try:
                    with open(asset_path, "wb") as f:
                        f.write(img_bytes)
                    block = ContentBlock(
                        block_type=BlockType.IMAGE,
                        source_format=SourceFormat.DOCX,
                        page_number=section_index,
                        reading_order=reading_order,
                        extraction_method=ExtractionMethod.NATIVE_IMAGE,
                        validation_status=ValidationStatus.VALID,
                        asset_id=asset_filename,
                        metadata={"r_embed": r_embed, "ext": ext},
                        notes=[
                            "Embedded image from DOCX. May contain a table, chart, or diagram. "
                            "Cannot be automatically identified without OCR."
                        ],
                    )
                    blocks.append(block)
                    reading_order += 1
                except Exception as exc:
                    logger.warning("Could not save image asset %s: %s", asset_filename, exc)
        return blocks

    # ── Table extraction ──────────────────────────────────────────

    def _extract_table(self, table: Table) -> TableData | None:
        rows: list[TableRow] = []
        col_count = 0

        for ri, row in enumerate(table.rows):
            is_header = (ri == 0)
            cells: list[TableCell] = []
            for ci, cell in enumerate(row.cells):
                text = cell.text.strip()
                cells.append(TableCell(
                    text=text,
                    row_index=ri,
                    col_index=ci,
                    is_header=is_header,
                ))
            col_count = max(col_count, len(cells))
            rows.append(TableRow(cells=cells, row_index=ri, is_header=is_header))

        if not rows:
            return None
        return TableData(
            rows=rows,
            column_count=col_count,
            row_count=len(rows),
            has_header=True,
        )

    def _validate_table(self, table_data: TableData) -> ValidationStatus:
        if table_data.row_count == 0 or table_data.column_count == 0:
            return ValidationStatus.FAILED
        non_empty = sum(
            1 for row in table_data.rows for cell in row.cells if cell.text.strip()
        )
        total = table_data.row_count * table_data.column_count
        if total > 0 and non_empty / total < 0.1:
            return ValidationStatus.SUSPECT
        if not table_data.html or "<table" not in table_data.html:
            return ValidationStatus.SUSPECT
        return ValidationStatus.VALID

    # ── Page-break detection ──────────────────────────────────────

    @staticmethod
    def _has_page_break(para_elem: Any) -> bool:
        for br in para_elem.iter(qn("w:br")):
            if br.get(qn("w:type")) in ("page", "column"):
                return True
        return False

    # ── Unsupported element check ─────────────────────────────────

    def _check_unsupported_elements(self, docx_path: str, result: DocumentResult) -> None:
        """
        Warn if the document contains text boxes or drawing shapes that
        python-docx cannot reliably extract.
        """
        try:
            doc = Document(docx_path)
            has_textbox = False
            has_shape   = False
            for para in doc.paragraphs:
                for run in para.runs:
                    xml = run._element.xml
                    if "txbx" in xml or "textBox" in xml:
                        has_textbox = True
                    if "wps:wsp" in xml or "mc:AlternateContent" in xml:
                        has_shape = True
            if has_textbox:
                result.summary.warnings.append(
                    "Document contains text boxes. Their content may not be fully extracted. "
                    "Limitation: python-docx does not expose text-box content in paragraph iteration."
                )
            if has_shape:
                result.summary.warnings.append(
                    "Document contains drawing shapes or grouped objects. "
                    "These are not reliably accessible via python-docx and may appear as images."
                )
        except Exception as exc:
            logger.debug("Unsupported element check failed: %s", exc)
